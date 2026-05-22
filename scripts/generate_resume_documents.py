#!/usr/bin/env python3
"""Generate ATS-friendly DOCX and PDF resume from structured content."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

WORKSPACE = Path(__file__).resolve().parent.parent
DOCX_OUT = WORKSPACE / "victor-akinla-software-engineer-resume.docx"
PDF_OUT = WORKSPACE / "victor-akinla-software-engineer-resume.pdf"
HTML_OUT = WORKSPACE / "victor-akinla-software-engineer-resume.html"

# Typography
FONT_NAME = "Calibri"
BODY_SIZE = Pt(10.5)
NAME_SIZE = Pt(22)
SECTION_SIZE = Pt(11)
SUBHEAD_SIZE = Pt(10.5)
MARGIN = Inches(0.65)
ACCENT = RGBColor(0x1A, 0x1A, 0x1A)
MUTED = RGBColor(0x44, 0x44, 0x44)


def strip_md(text: str) -> str:
    return re.sub(r"\*\*(.+?)\*\*", r"\1", text)


def add_run_bold(paragraph, text: str, bold_parts: list[str] | None = None, size=None, color=None):
    """Add text with optional **bold** segments from markdown-style markers."""
    if bold_parts is None:
        parts = re.split(r"(\*\*.+?\*\*)", text)
        for part in parts:
            m = re.match(r"\*\*(.+?)\*\*", part)
            run = paragraph.add_run(m.group(1) if m else part)
            run.bold = bool(m)
            run.font.name = FONT_NAME
            if size:
                run.font.size = size
            if color:
                run.font.color.rgb = color
        return
    run = paragraph.add_run(text)
    run.font.name = FONT_NAME
    if size:
        run.font.size = size


def set_document_defaults(doc: Document) -> None:
    for section in doc.sections:
        section.top_margin = MARGIN
        section.bottom_margin = MARGIN
        section.left_margin = MARGIN
        section.right_margin = MARGIN

    style = doc.styles["Normal"]
    style.font.name = FONT_NAME
    style.font.size = BODY_SIZE
    style.paragraph_format.space_after = Pt(2)
    style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    style.paragraph_format.line_spacing = 1.08


def add_section_heading(doc: Document, title: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(title.upper())
    run.bold = True
    run.font.name = FONT_NAME
    run.font.size = SECTION_SIZE
    run.font.color.rgb = ACCENT
    # Bottom border via paragraph border XML
    pPr = p._p.get_or_add_pPr()
    pBdr = pPr.find(qn("w:pBdr"))
    if pBdr is None:
        from docx.oxml import OxmlElement

        pBdr = OxmlElement("w:pBdr")
        pPr.append(pBdr)
    from docx.oxml import OxmlElement

    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "1A1A1A")
    pBdr.append(bottom)


def add_contact(doc: Document) -> None:
    name = doc.add_paragraph()
    name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    nr = name.add_run("Victor Akinla")
    nr.bold = True
    nr.font.name = FONT_NAME
    nr.font.size = NAME_SIZE
    nr.font.color.rgb = ACCENT

    contact = doc.add_paragraph()
    contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cr = contact.add_run(
        "Dublin, Ireland  |  Irish  |  +353 86 073 1698  |  vakinla.va@gmail.com"
    )
    cr.font.name = FONT_NAME
    cr.font.size = Pt(10)
    cr.font.color.rgb = MUTED

    links = doc.add_paragraph()
    links.alignment = WD_ALIGN_PARAGRAPH.CENTER
    lr = links.add_run(
        "LinkedIn: linkedin.com/in/victor-akinla-21061a88  |  GitHub: github.com/akavic"
    )
    lr.font.name = FONT_NAME
    lr.font.size = Pt(9.5)
    lr.font.color.rgb = MUTED


def add_body_paragraph(doc: Document, text: str, space_after: int = 4) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    add_run_bold(p, text, size=BODY_SIZE)


def add_skill_line(doc: Document, label: str, value: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.left_indent = Inches(0)
    lr = p.add_run(f"{label}: ")
    lr.bold = True
    lr.font.name = FONT_NAME
    lr.font.size = BODY_SIZE
    vr = p.add_run(value)
    vr.font.name = FONT_NAME
    vr.font.size = BODY_SIZE


def add_job_header(doc: Document, company: str, role: str, dates: str, location: str = "") -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(0)
    c = p.add_run(f"{company} — {role}")
    c.bold = True
    c.font.name = FONT_NAME
    c.font.size = SUBHEAD_SIZE

    meta = doc.add_paragraph()
    meta.paragraph_format.space_after = Pt(3)
    mr = meta.add_run(f"{dates}" + (f"  |  {location}" if location else ""))
    mr.italic = True
    mr.font.name = FONT_NAME
    mr.font.size = Pt(10)
    mr.font.color.rgb = MUTED


def add_subrole(doc: Document, title: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(title)
    r.bold = True
    r.font.name = FONT_NAME
    r.font.size = BODY_SIZE


def add_bullet(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.left_indent = Inches(0.2)
    add_run_bold(p, text, size=BODY_SIZE)


def build_docx() -> Document:
    doc = Document()
    set_document_defaults(doc)
    add_contact(doc)

    add_section_heading(doc, "Professional Summary")
    add_body_paragraph(
        doc,
        "Software Engineer with production **full-stack** and **backend** experience delivering features across **React**, **Go**, **PHP**, **PostgreSQL**, **REST APIs**, and **microservices**—including **API integration**, **transactional processing**, **authentication**, **GitLab CI/CD** deployments, and **cross-stack debugging**.",
    )
    add_body_paragraph(
        doc,
        "Enterprise **identity management** background (SailPoint, Oracle Identity Analytics) with **Agile** delivery; targeting Software Engineer, Full-Stack, Backend, Platform, API, React, and Golang roles.",
        space_after=2,
    )

    add_section_heading(doc, "Technical Skills")
    skills = {
        "Languages": "JavaScript, TypeScript, Go (Golang), PHP, SQL, HTML, CSS, Shell scripting",
        "Frontend": "React 17, Redux, state management, custom hooks, form validation, REST API consumption",
        "Backend": "Go REST APIs, GORM, PHP, Node.js, RESTful API design, CRUD, transactional processing, API integration, authentication, microservices",
        "Databases": "PostgreSQL, relational modeling, multi-table transactional writes, schema alignment, Oracle SQL",
        "Cloud & DevOps": "GitLab CI/CD, Git, cross-compilation (Go/Linux), Apache, staging and production deployments, cloud deployment workflows, Azure fundamentals (AZ-900)",
        "Tools & Platforms": "Linux, Unix, Agile/Scrum, Jira, Eclipse, XML, SailPoint, Oracle Identity Analytics (OIA), IDAM, ITIL, BMC Remedy, distributed systems, production support",
    }
    for label, value in skills.items():
        add_skill_line(doc, label, value)

    add_section_heading(doc, "Professional Experience")

    add_job_header(doc, "WiLine", "Software Developer", "August 2022 – Present", "Dublin, Ireland")
    add_body_paragraph(
        doc,
        "Engineer on customer and internal platforms combining **React** frontends, **Go REST APIs**, **PHP** legacy services, and **PostgreSQL** under a **microservices** architecture. Owns feature delivery, **backend development**, **API integration**, production defect resolution, and **GitLab CI/CD** deployments.",
        space_after=3,
    )
    wiline_bullets = [
        "Delivered an enterprise **Vendor Contract Management System** end to end across **React**, **Go**, and **PHP**: dynamic multi-type contract workflows (Type II, Real Estate, and others) with tailored validation, reusable hooks, and **state management** via custom hooks and **Redux**.",
        "Implemented production **RESTful APIs** in **Go** for full **CRUD** on vendor contracts, including **transactional processing** across related **PostgreSQL** entities (vendor contacts, NOC contacts, addresses) with **GORM** and strict struct/schema alignment.",
        "Built secure **server-to-server integration** between Go services and legacy **PHP** contract file storage using shared **API keys** and **Base64 Basic Authentication**; added standalone PHP endpoints for external upload/download without PHP session dependency.",
        "Resolved complex **cross-stack production bugs** spanning JavaScript, Go/database field mismatches, JSON unmarshal errors, and API edge cases—restoring stability for live users.",
        "Operated **GitLab CI/CD** pipelines for test and production: React builds, **Go binary** cross-compilation for Linux, and PHP releases—supporting repeatable **cloud deployment** practices.",
        "Developed and maintained **full-stack** web applications across **PHP**, **React**, **JavaScript**, and **PostgreSQL**; shipped enhancements and performance fixes through structured **debugging** and verification.",
        "Designed, built, and deployed standalone **Node.js**, **React**, and **PostgreSQL** operational tools to production and staging with ownership of business logic and data access.",
        "Partnered with cross-functional stakeholders in **Agile development** to plan work, unblock integrations, and support **production support** for released features.",
    ]
    for b in wiline_bullets:
        add_bullet(doc, b)

    add_job_header(doc, "Wipro", "Application Engineer", "April 2019 – July 2022", "Dublin, Ireland")
    add_body_paragraph(
        doc,
        "Application engineer for global **identity access management** programs, focused on defect resolution, integration testing, and reliable enterprise releases.",
        space_after=3,
    )
    add_subrole(doc, "SailPoint | 2020 – 2022")
    for b in [
        "Diagnosed and fixed application defects with development teams via **Jira**, delivering **testable** code changes validated through **unit**, **integration**, and **UAT** testing.",
        "Supported client-facing **UAT** and integration test cycles with engineers and customers to meet release quality bars.",
        "Applied **Git**, **Eclipse**, and **XML** in development environments to implement and verify fixes on IAM application codebases.",
        "Authored defect resolution documentation for engineering and support knowledge reuse.",
        "Raised **ITIL**-aligned change requests through **BMC Remedy** for controlled production change.",
        "Led **batch run** coordination using **Unix shell scripting**, PuTTY, and WinSCP across dependent teams.",
    ]:
        add_bullet(doc, b)

    add_subrole(doc, "Oracle Identity Analytics (OIA) | 2019 – 2020")
    for b in [
        "Integrated new resources into OIA using **Oracle SQL** and cross-team delivery for identity management initiatives.",
        "Executed **server-side testing** after patches—confirming application, web, and NFS/logging health on Linux servers before release approval.",
    ]:
        add_bullet(doc, b)

    add_job_header(doc, "Host Ireland", "Network Engineer", "September 2018 – March 2019")
    for b in [
        "Provisioned customer IP subnets and resolved network incidents through help-desk support.",
        "Configured customer routing paths on **Juniper** network devices.",
    ]:
        add_bullet(doc, b)

    add_section_heading(doc, "Projects")
    add_job_header(doc, "WiLine", "Customer Satisfaction Voting System", "December 2022 – February 2023")
    for b in [
        "Built a **React** web application with **PostgreSQL** persistence to collect customer satisfaction votes and feedback for service improvement.",
        "Owned front-end implementation and database-backed storage for reliable feedback capture in production.",
    ]:
        add_bullet(doc, b)

    add_job_header(doc, "WiLine", "Operational Automation Utilities", "2022 – Present")
    add_bullet(
        doc,
        "Created standalone **Node.js**, **React**, and **PostgreSQL** scripts for internal workflows including data handling, application logic, and deployment to staging/production environments.",
    )

    add_section_heading(doc, "Education")
    p = doc.add_paragraph()
    r1 = p.add_run("BSc Computer Science — Dublin City University")
    r1.bold = True
    r1.font.name = FONT_NAME
    r1.font.size = BODY_SIZE
    meta = doc.add_paragraph()
    m = meta.add_run("August 2013 – May 2017  |  Dublin, Ireland")
    m.font.name = FONT_NAME
    m.font.size = Pt(10)
    m.font.color.rgb = MUTED
    add_body_paragraph(doc, "Relevant study: data structures and algorithms, databases, cryptography.", space_after=2)

    add_section_heading(doc, "Certifications")
    for cert in [
        "Microsoft Certified: Azure Fundamentals (AZ-900)",
        "TOGAF 9 Certified — The Open Group",
    ]:
        add_bullet(doc, cert)

    add_section_heading(doc, "Additional Technical Highlights")
    for b in [
        "**Full-stack ownership:** UI (React/Redux), API layers (Go/PHP), and PostgreSQL data models in the same delivery cycle.",
        "**Legacy + modern integration:** Bridging new **Go** services with established **PHP** platforms without disrupting existing workflows.",
        "**Production engineering:** Deployment discipline, integration testing, and live defect triage across distributed application tiers.",
        "**Enterprise IAM exposure:** SailPoint and OIA experience with security-sensitive systems, change control, and auditability.",
    ]:
        add_bullet(doc, b)

    return doc


def build_html() -> str:
    """ATS-friendly single-column HTML for PDF export."""
    summary = (
        "<p>Software Engineer with production <strong>full-stack</strong> and <strong>backend</strong> "
        "experience delivering features across <strong>React</strong>, <strong>Go</strong>, <strong>PHP</strong>, "
        "<strong>PostgreSQL</strong>, <strong>REST APIs</strong>, and <strong>microservices</strong>—including "
        "<strong>API integration</strong>, <strong>transactional processing</strong>, <strong>authentication</strong>, "
        "<strong>GitLab CI/CD</strong> deployments, and <strong>cross-stack debugging</strong>.</p>"
        "<p>Enterprise <strong>identity management</strong> background (SailPoint, Oracle Identity Analytics) with "
        "<strong>Agile</strong> delivery; targeting Software Engineer, Full-Stack, Backend, Platform, API, React, "
        "and Golang roles.</p>"
    )

    skills_html = """
    <p><span class="label">Languages:</span> JavaScript, TypeScript, Go (Golang), PHP, SQL, HTML, CSS, Shell scripting</p>
    <p><span class="label">Frontend:</span> React 17, Redux, state management, custom hooks, form validation, REST API consumption</p>
    <p><span class="label">Backend:</span> Go REST APIs, GORM, PHP, Node.js, RESTful API design, CRUD, transactional processing, API integration, authentication, microservices</p>
    <p><span class="label">Databases:</span> PostgreSQL, relational modeling, multi-table transactional writes, schema alignment, Oracle SQL</p>
    <p><span class="label">Cloud &amp; DevOps:</span> GitLab CI/CD, Git, cross-compilation (Go/Linux), Apache, staging and production deployments, cloud deployment workflows, Azure fundamentals (AZ-900)</p>
    <p><span class="label">Tools &amp; Platforms:</span> Linux, Unix, Agile/Scrum, Jira, Eclipse, XML, SailPoint, OIA, IDAM, ITIL, BMC Remedy, distributed systems, production support</p>
    """

    def bullets(items: list[str]) -> str:
        return "<ul>" + "".join(f"<li>{strip_md(i)}</li>" for i in items) + "</ul>"

    wiline = bullets(
        [
            "Delivered an enterprise Vendor Contract Management System end to end across React, Go, and PHP: dynamic multi-type contract workflows (Type II, Real Estate, and others) with tailored validation, reusable hooks, and state management via custom hooks and Redux.",
            "Implemented production RESTful APIs in Go for full CRUD on vendor contracts, including transactional processing across related PostgreSQL entities (vendor contacts, NOC contacts, addresses) with GORM and strict struct/schema alignment.",
            "Built secure server-to-server integration between Go services and legacy PHP contract file storage using shared API keys and Base64 Basic Authentication; added standalone PHP endpoints for external upload/download without PHP session dependency.",
            "Resolved complex cross-stack production bugs spanning JavaScript, Go/database field mismatches, JSON unmarshal errors, and API edge cases—restoring stability for live users.",
            "Operated GitLab CI/CD pipelines for test and production: React builds, Go binary cross-compilation for Linux, and PHP releases—supporting repeatable cloud deployment practices.",
            "Developed and maintained full-stack web applications across PHP, React, JavaScript, and PostgreSQL; shipped enhancements and performance fixes through structured debugging and verification.",
            "Designed, built, and deployed standalone Node.js, React, and PostgreSQL operational tools to production and staging with ownership of business logic and data access.",
            "Partnered with cross-functional stakeholders in Agile development to plan work, unblock integrations, and support production support for released features.",
        ]
    )

    return f"""<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="utf-8"/>
<title>Victor Akinla — Software Engineer Resume</title>
<style>
  @page {{ size: A4; margin: 0.55in 0.6in; }}
  body {{
    font-family: Calibri, Arial, Helvetica, sans-serif;
    font-size: 10.5pt;
    line-height: 1.25;
    color: #1a1a1a;
    max-width: 7.2in;
    margin: 0 auto;
  }}
  h1 {{
    font-size: 22pt;
    text-align: center;
    margin: 0 0 4px 0;
    letter-spacing: 0.3px;
  }}
  .contact, .links {{
    text-align: center;
    color: #444;
    margin: 0 0 2px 0;
    font-size: 10pt;
  }}
  .links {{ font-size: 9.5pt; margin-bottom: 10px; }}
  h2 {{
    font-size: 11pt;
    text-transform: uppercase;
    letter-spacing: 0.5px;
    border-bottom: 1px solid #1a1a1a;
    padding-bottom: 2px;
    margin: 12px 0 6px 0;
  }}
  .job-title {{ font-weight: bold; font-size: 10.5pt; margin: 8px 0 0 0; }}
  .job-meta {{ font-style: italic; color: #444; font-size: 10pt; margin: 0 0 4px 0; }}
  .subrole {{ font-weight: bold; margin: 6px 0 2px 0; }}
  p {{ margin: 0 0 4px 0; }}
  ul {{ margin: 2px 0 6px 0; padding-left: 18px; }}
  li {{ margin-bottom: 2px; }}
  .label {{ font-weight: bold; }}
</style>
</head>
<body>
<h1>Victor Akinla</h1>
<p class="contact">Dublin, Ireland &nbsp;|&nbsp; Irish &nbsp;|&nbsp; +353 86 073 1698 &nbsp;|&nbsp; vakinla.va@gmail.com</p>
<p class="links">LinkedIn: linkedin.com/in/victor-akinla-21061a88 &nbsp;|&nbsp; GitHub: github.com/akavic</p>

<h2>Professional Summary</h2>
{summary}

<h2>Technical Skills</h2>
{skills_html}

<h2>Professional Experience</h2>
<p class="job-title">WiLine — Software Developer</p>
<p class="job-meta">August 2022 – Present &nbsp;|&nbsp; Dublin, Ireland</p>
<p>Engineer on customer and internal platforms combining React frontends, Go REST APIs, PHP legacy services, and PostgreSQL under a microservices architecture.</p>
{wiline}

<p class="job-title">Wipro — Application Engineer</p>
<p class="job-meta">April 2019 – July 2022 &nbsp;|&nbsp; Dublin, Ireland</p>
<p class="subrole">SailPoint | 2020 – 2022</p>
{bullets([
"Diagnosed and fixed application defects with development teams via Jira, delivering testable code changes validated through unit, integration, and UAT testing.",
"Supported client-facing UAT and integration test cycles with engineers and customers to meet release quality bars.",
"Applied Git, Eclipse, and XML in development environments to implement and verify fixes on IAM application codebases.",
"Authored defect resolution documentation for engineering and support knowledge reuse.",
"Raised ITIL-aligned change requests through BMC Remedy for controlled production change.",
"Led batch run coordination using Unix shell scripting, PuTTY, and WinSCP across dependent teams.",
])}
<p class="subrole">Oracle Identity Analytics (OIA) | 2019 – 2020</p>
{bullets([
"Integrated new resources into OIA using Oracle SQL and cross-team delivery for identity management initiatives.",
"Executed server-side testing after patches—confirming application, web, and NFS/logging health on Linux servers before release approval.",
])}

<p class="job-title">Host Ireland — Network Engineer</p>
<p class="job-meta">September 2018 – March 2019</p>
{bullets([
"Provisioned customer IP subnets and resolved network incidents through help-desk support.",
"Configured customer routing paths on Juniper network devices.",
])}

<h2>Projects</h2>
<p class="job-title">Customer Satisfaction Voting System — WiLine</p>
<p class="job-meta">December 2022 – February 2023</p>
{bullets([
"Built a React web application with PostgreSQL persistence to collect customer satisfaction votes and feedback for service improvement.",
"Owned front-end implementation and database-backed storage for reliable feedback capture in production.",
])}
<p class="job-title">Operational Automation Utilities — WiLine</p>
<p class="job-meta">2022 – Present</p>
{bullets([
"Created standalone Node.js, React, and PostgreSQL scripts for internal workflows including data handling, application logic, and deployment to staging/production environments.",
])}

<h2>Education</h2>
<p><strong>BSc Computer Science — Dublin City University</strong><br/>
<span class="job-meta">August 2013 – May 2017 &nbsp;|&nbsp; Dublin, Ireland</span></p>
<p>Relevant study: data structures and algorithms, databases, cryptography.</p>

<h2>Certifications</h2>
{bullets([
"Microsoft Certified: Azure Fundamentals (AZ-900)",
"TOGAF 9 Certified — The Open Group",
])}

<h2>Additional Technical Highlights</h2>
{bullets([
"Full-stack ownership across UI (React/Redux), API layers (Go/PHP), and PostgreSQL data models.",
"Legacy + modern integration bridging Go services with established PHP platforms.",
"Production engineering: deployment discipline, integration testing, and live defect triage.",
"Enterprise IAM exposure with SailPoint and OIA in security-sensitive environments.",
])}
</body>
</html>"""


def export_pdf(html: str) -> bool:
    try:
        from weasyprint import HTML

        HTML(string=html).write_pdf(str(PDF_OUT))
        return True
    except Exception as exc:
        print(f"PDF export skipped: {exc}")
        return False


def main() -> None:
    doc = build_docx()
    doc.save(DOCX_OUT)
    print(f"DOCX: {DOCX_OUT}")

    html = build_html()
    HTML_OUT.write_text(html, encoding="utf-8")
    print(f"HTML: {HTML_OUT}")

    if export_pdf(html):
        print(f"PDF:  {PDF_OUT}")


if __name__ == "__main__":
    main()
